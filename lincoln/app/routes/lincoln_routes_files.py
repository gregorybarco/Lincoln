"""
Lincoln File Routes  v0.6.0
=============================
Flask routes for file attachment handling.

Changes in v0.6.0:
  - Size limits read from DB settings (upload_max_text_kb, upload_max_doc_mb)
  - Extended language support: Julia, R, MATLAB, Maple, SAS, Stata, GAMS, AMPL, etc.
  - Image file support: .png .jpg .jpeg .gif .bmp .tiff .webp
    -> Tesseract OCR extraction (text images)
    -> Vision model extraction (charts/surfaces) via ?mode=vision
  - Maple .mw worksheet extraction (XML-based)
  - GET /api/files/browse -- folder browser for sidebar file picker
  - BUG FIX B4: file path handling returns File object, not re-opens picker

Endpoints:
  POST /api/files/upload          Upload and extract a file for chat injection
  GET  /api/files/browse          Browse a folder path (sidebar file picker)
"""

import hashlib
import os
import uuid
import xml.etree.ElementTree as ET
from pathlib import Path

from flask import Blueprint, jsonify, request

from lincoln.lincoln_configuration import DB_PATH
from lincoln.lincoln_database import get_setting

files_blueprint = Blueprint("files", __name__)

_UPLOAD_DIR = DB_PATH.parent / "uploads"
_UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

# ── Extension categories ──────────────────────────────────────────────────────

_TEXT_EXTENSIONS = {
    # Python
    ".py", ".pyi",
    # Fortran
    ".f90", ".f", ".for", ".f95", ".f03", ".f08", ".fpp",
    # C / C++
    ".c", ".cc", ".cpp", ".cxx", ".h", ".hh", ".hpp", ".hxx",
    # Web
    ".html", ".htm", ".css", ".scss", ".sass", ".less",
    ".js", ".mjs", ".cjs", ".jsx", ".ts", ".tsx", ".svelte", ".vue",
    # Shell / script
    ".sh", ".bash", ".zsh", ".ps1", ".bat", ".cmd",
    # Config / data
    ".toml", ".yaml", ".yml", ".ini", ".cfg", ".conf",
    ".json", ".xml", ".gradle", ".cmake",
    # Docs
    ".md", ".rst", ".txt",
    # SQL
    ".sql",
    # Other languages
    ".java", ".kt", ".scala", ".cs", ".fs", ".go", ".rs",
    ".swift", ".rb", ".php", ".pl", ".lua",
    # Julia
    ".jl",
    # MATLAB / Octave
    ".m",
    # Mathematica / Wolfram
    ".nb", ".wl",
    # R
    ".r", ".Rmd", ".rmd",
    # Stats / econometrics
    ".sas", ".do", ".ado",
    # Optimisation
    ".gms", ".ampl", ".lp", ".mps",
    # Maple (text/procedure files)
    ".mpl", ".maple", ".mm",
    # LaTeX
    ".tex", ".latex", ".bib",
    # Jupyter
    ".ipynb",
}

# Document types requiring binary extraction
_DOC_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".csv", ".mw"}

# Image types (OCR or vision model extraction)
_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff", ".tif", ".webp"}


# ── Size limits (from DB, defaulting to .env-era values) ─────────────────────

def _max_text_bytes() -> int:
    try:
        return int(get_setting("upload_max_text_kb", "512")) * 1024
    except Exception:
        return 512 * 1024


def _max_doc_bytes() -> int:
    try:
        return int(get_setting("upload_max_doc_mb", "2")) * 1024 * 1024
    except Exception:
        return 2 * 1024 * 1024


# ── Maple .mw worksheet extraction ───────────────────────────────────────────

def _extract_mw(raw: bytes) -> str:
    """
    Extract readable text from a Maple .mw worksheet (XML format).
    Maple worksheets are well-formed XML; we extract Text and Input elements.
    """
    try:
        root  = ET.fromstring(raw.decode("utf-8", errors="replace"))
        parts = []
        for elem in root.iter():
            # Input cells (Maple code)
            if "Input" in (elem.tag or ""):
                text = "".join(elem.itertext()).strip()
                if text:
                    parts.append(f"[Maple Input]\n{text}")
            # Text cells (prose / annotations)
            elif "Text" in (elem.tag or "") and elem.text:
                parts.append(elem.text.strip())
        return "\n\n".join(parts) if parts else "(Maple worksheet: no readable content extracted)"
    except ET.ParseError:
        # Some .mw files use a binary header; fall back to raw text extraction
        try:
            text = raw.decode("utf-8", errors="replace")
            return text[:8192]
        except Exception:
            return "(Maple worksheet: could not parse XML)"
    except Exception as exc:
        return f"(Maple worksheet extraction failed: {exc})"


# ── Main upload handler ───────────────────────────────────────────────────────

@files_blueprint.route("/api/files/upload", methods=["POST"])
def upload_file():
    """
    Accept a file upload, extract its text content, save to uploads/.
    Returns a file_id the client uses when sending the chat message.

    Query params:
      mode         : 'ocr' (default) or 'vision' for image files
      vision_model : model name to use for vision extraction
      lang         : Tesseract language code (default 'eng')
    """
    if "file" not in request.files:
        return jsonify({"error": "No file in request"}), 400

    file      = request.files["file"]
    filename  = file.filename or "upload"
    extension = Path(filename).suffix.lower()
    raw       = file.read()

    # ── Route by extension ────────────────────────────────────────────────────

    if extension in _TEXT_EXTENSIONS:
        if len(raw) > _max_text_bytes():
            return jsonify({
                "error": (
                    f"File too large. Maximum for text/code files is "
                    f"{get_setting('upload_max_text_kb', '512')} KB. "
                    f"Increase this limit in Settings -> Uploads."
                )
            }), 413

        try:
            extracted = raw.decode("utf-8", errors="replace")
        except Exception as exc:
            return jsonify({"error": f"Could not decode text file: {exc}"}), 400

    elif extension in _IMAGE_EXTENSIONS:
        mode         = request.args.get("mode", "ocr")
        vision_model = request.args.get("vision_model", "")
        lang         = request.args.get("lang", "eng")

        from lincoln.lincoln_ocr_service import extract_image
        from lincoln.lincoln_configuration import OLLAMA_BASE_URL

        extracted = extract_image(
            image_bytes   = raw,
            mode          = mode,
            lang          = lang,
            psm           = 6,
            vision_model  = vision_model,
            ollama_url    = OLLAMA_BASE_URL,
        )

    elif extension == ".mw":
        if len(raw) > _max_doc_bytes():
            return jsonify({"error": "Maple worksheet too large"}), 413
        extracted = _extract_mw(raw)

    elif extension == ".pdf":
        if len(raw) > _max_doc_bytes():
            return jsonify({"error": f"PDF too large. Maximum is {get_setting('upload_max_doc_mb', '2')} MB."}), 413
        try:
            import fitz  # PyMuPDF
            doc       = fitz.open(stream=raw, filetype="pdf")
            pages     = [page.get_text() for page in doc]
            extracted = "\n\n".join(pages)
        except ImportError:
            return jsonify({
                "error": "PDF support requires PyMuPDF. Run: pip install pymupdf"
            }), 500
        except Exception as exc:
            return jsonify({"error": f"PDF extraction failed: {exc}"}), 500

    elif extension == ".docx":
        if len(raw) > _max_doc_bytes():
            return jsonify({"error": "Document too large"}), 413
        try:
            import io
            from docx import Document
            doc       = Document(io.BytesIO(raw))
            extracted = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except ImportError:
            return jsonify({
                "error": "Word document support requires python-docx. Run: pip install python-docx"
            }), 500
        except Exception as exc:
            return jsonify({"error": f"Word document extraction failed: {exc}"}), 500

    elif extension == ".xlsx":
        if len(raw) > _max_doc_bytes():
            return jsonify({"error": "Spreadsheet too large"}), 413
        try:
            import io
            import openpyxl
            wb   = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
            rows = []
            for sheet in wb.worksheets:
                rows.append(f"[Sheet: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    row_text = "\t".join(
                        str(cell) if cell is not None else "" for cell in row
                    )
                    if row_text.strip():
                        rows.append(row_text)
            extracted = "\n".join(rows)
        except ImportError:
            return jsonify({
                "error": "Excel support requires openpyxl. Run: pip install openpyxl"
            }), 500
        except Exception as exc:
            return jsonify({"error": f"Excel extraction failed: {exc}"}), 500

    elif extension == ".csv":
        if len(raw) > _max_text_bytes():
            return jsonify({"error": "CSV too large"}), 413
        extracted = raw.decode("utf-8", errors="replace")

    else:
        return jsonify({
            "error": (
                f"File type '{extension}' is not supported for text extraction. "
                f"Supported types include: Python, Fortran, C/C++, Julia, R, MATLAB, "
                f"Maple, LaTeX, Markdown, PDF, Word, Excel, CSV, and most code files."
            )
        }), 415

    # ── Save extracted text ───────────────────────────────────────────────────

    file_id   = uuid.uuid4().hex
    save_path = _UPLOAD_DIR / f"{file_id}.txt"
    save_path.write_text(extracted, encoding="utf-8")

    return jsonify({
        "file_id":   file_id,
        "filename":  filename,
        "extension": extension,
        "size":      len(extracted),
        "preview":   extracted[:200] + ("..." if len(extracted) > 200 else ""),
    })


# ── Folder browser (sidebar file picker) ──────────────────────────────────────

@files_blueprint.route("/api/files/browse", methods=["GET"])
def browse_folder():
    """
    Return the contents of a directory for the sidebar file browser.
    Used by the file attach picker so the user can navigate and select files
    without the browser's native file picker opening.

    Query params:
      path : absolute directory path to list (default: user home)
    """
    raw_path = request.args.get("path", "").strip()

    if not raw_path:
        import os
        raw_path = os.path.expanduser("~")

    target = Path(raw_path)

    if not target.exists() or not target.is_dir():
        return jsonify({"error": f"Path does not exist or is not a directory: {raw_path}"}), 400

    try:
        entries = []
        for entry in sorted(target.iterdir(), key=lambda e: (not e.is_dir(), e.name.lower())):
            try:
                ext   = entry.suffix.lower() if entry.is_file() else ""
                is_attachable = (
                    ext in _TEXT_EXTENSIONS
                    or ext in _DOC_EXTENSIONS
                    or ext in _IMAGE_EXTENSIONS
                )
                entries.append({
                    "name":         entry.name,
                    "path":         str(entry),
                    "is_dir":       entry.is_dir(),
                    "extension":    ext,
                    "is_attachable": is_attachable,
                    "size":         entry.stat().st_size if entry.is_file() else None,
                })
            except OSError:
                continue

        return jsonify({
            "path":    str(target),
            "parent":  str(target.parent),
            "entries": entries,
        })

    except PermissionError:
        return jsonify({"error": f"Permission denied: {raw_path}"}), 403
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500
